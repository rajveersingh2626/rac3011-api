import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('backlogs_data/Club Project Submissions - July & August 2026 (Responses).xlsx', data_only=True)
sheet = wb.active

authors = [
    'Sania Khatter', 'Janhvee Sharma', 'Nitika Khatri', 'Anshita Jain', 'mishti bansal',
    'Srishty Goyal', 'Rashika Yadav', 'Harshita Gupta', 'Leysha Tewatia', 'Mohd Arham',
    'Aryan Sanjeev', 'Krishika Bindal', 'Aryan Yadav', 'Rishika Ranjan', 'Yashika Kapoor',
    'Mansi Mishra', 'Galgotia', 'NSIT', 'yogya goyal', 'Gauri Sharma', 'ajay chawla',
    'Ashutosh Bhandari', 'Rishabh Babber', 'UTKARSH DUTTA', 'Ananya', 'somya', 'Sonal Dhingra'
]

print(f"Total rows in Excel: {sheet.max_row}")

# Check all cell contents in each row to see which club / project each author belongs to
for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
    row_text = " ".join([str(c) for c in row if c is not None]).lower()
    matched = []
    for a in authors:
        if a.lower() in row_text:
            matched.append(a)
    club = row[1]
    p1_name = row[2]
    p2_name = row[12] if len(row) > 12 else ''
    if matched:
        print(f"Row {row_idx}: Club='{club}', P1='{p1_name}', P2='{p2_name}', Matched={matched}")
    else:
        print(f"Row {row_idx}: Club='{club}', P1='{p1_name}', P2='{p2_name}' [NO MATCH]")
